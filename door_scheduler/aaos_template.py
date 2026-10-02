"""Map complete door schedule columns onto country AAOS import templates."""

from __future__ import annotations

import re
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

TEMPLATE_ROOT = Path(__file__).resolve().parent.parent / "Door_Schedule_Template"

COUNTRY_NAMES: dict[str, str] = {
    "CH": "China",
    "HK": "Hong Kong",
    "ID": "Indonesia",
    "MY": "Malaysia",
    "PH": "Philippines",
    "SG": "Singapore",
    "TH": "Thailand",
    "TW": "Taiwan",
    "VN": "Vietnam",
}

# Logical fields → possible AAOS Schedule header labels across countries.
AAOS_FIELD_ALIASES: dict[str, list[str]] = {
    "mark": ["Mark", "门号", "標記"],
    "building": ["Building", "楼号", "大樓"],
    "usage": ["Usage", "使用频率", "使用方式"],
    "level": ["Level", "楼层", "等級"],
    "level_marks": ["Level Marks"],
    "room": ["Room", "房间内部描述"],
    "from_room": ["From Room", "房间外部描述", "從房間"],
    "to_room": ["To Room", "到房間", "房间内部描述"],
    "exterior": ["Exterior", "是否外门", "外觀"],
    "qty": ["Qty", "数量", "數量"],
    "door_location": ["Door Location", "门位置", "門位置"],
    "config": [
        "Config",
        "Configuration",
        "开启方式",
        "配置",
    ],
    "width": ["Width", "洞口宽度（mm）", "洞口宽度(mm)", "寬度"],
    "leaf1_width": ["Leaf 1 Width", "Leaf1 Width", "门扇1宽度", "門扇 1 寬度"],
    "leaf2_width": ["Leaf 2 Width", "Leaf2 Width", "门扇2宽度", "門扇 2 寬度"],
    "height": ["Height", "洞口高度（mm）", "洞口高度(mm)", "高度"],
    "thickness": ["thickness", "Thickness", "门厚（mm）", "门厚(mm)", "厚度"],
    "handing": ["handing", "Handing", "手向", "門扇開關方向"],
    "door_type": ["door type", "Door Type", "门型", "門類型"],
    "door_material": ["door material", "Door Material", "门扇材质", "門材質"],
    "door_elevation": ["door Elevation", "Door Elevation"],
    "frame_material": ["frame material", "Frame Material", "门框材质", "門框材質"],
    "frame_elevation": ["frame Elevation", "Frame Elevation"],
    "frame_type": ["Frame Type", "门框类型", "門框類型"],
    "fire_rating": ["fire rating", "Fire Rating", "防火等级", "防火等級"],
    "smoke_rating": ["Smoke Rating", "防烟等级", "隔煙等級"],
    "acoustic_rating": ["Acoustic Rating", "隔声等级", "隔音等級"],
    "hwset": ["hwset", "Hardware", "五金组别", "五金"],
    "door_glazing": ["Door Glazing"],
    "door_glazing_thickness": ["Door Glazing (Thickness)"],
    "door_glazing_vision": ["Door Glazing and Vision Panel"],
    "comments": ["Comments", "备注", "註解"],
    "specifier_remarks": [
        "specifier remarks",
        "Specifier Remarks",
        "五金顾问备注",
        "指定者備註",
    ],
    "note1": ["Note 1", "备注2", "備註 1"],
    "note2": ["Note 2", "备注3", "備註 2"],
    "note3": ["Note 3", "備註 3"],
    "note4": ["Note 4"],
    "rfi": ["RFI"],
    "depth": ["Depth", "深度"],
    "arch_door_material": ["Arch Door Material", "建筑门扇材质", "拱門材質"],
    "arch_door_type": ["Arch Door Type", "建筑门型", "拱門類型"],
    "arch_frame_material": ["Arch Frame Material", "建筑门框材质", "拱門框材質"],
    "arch_frame_type": ["Arch Frame Type", "建筑框型", "拱門框類型"],
    "arch_door_finish": ["Arch Door Finish"],
    "arch_frame_finish": ["Arch Frame Finish"],
    "fitout_door_model": ["精装门型号"],
}

# Logical field → preferred complete-schedule source columns.
SOURCE_FOR_FIELD: dict[str, list[str]] = {
    "mark": ["Door Mark", "Door Type", "Mark"],
    "building": ["Building"],
    "usage": ["Usage"],
    "level": ["Level"],
    "level_marks": ["Level Marks"],
    "room": ["Location (To Room)", "To Room", "Location", "Room"],
    "from_room": ["Location (From Room)", "From Room"],
    "to_room": ["Location (To Room)", "To Room", "Location"],
    "exterior": ["Exterior"],
    "qty": ["Qty", "Quantity"],
    "door_location": ["Door Location", "Location"],
    "config": ["Configuration of Door", "Configuration", "Config"],
    "width": ["Width (mm)", "Width"],
    "leaf1_width": ["Leaf 1 Width", "Leaf1 Width"],
    "leaf2_width": ["Leaf 2 Width", "Leaf2 Width"],
    "height": ["Height (mm)", "Height"],
    "thickness": ["Thickness (mm)", "Thickness", "thickness"],
    "handing": ["Handing", "handing"],
    "door_type": ["Door Type", "door type"],
    "door_material": ["Door Material", "door material"],
    "door_elevation": ["door Elevation", "Door Elevation"],
    "frame_material": ["Frame Material", "frame material"],
    "frame_elevation": ["frame Elevation", "Frame Elevation"],
    "frame_type": ["Frame Type"],
    "fire_rating": ["Fire Rating", "fire rating"],
    "smoke_rating": ["Smoke Rating"],
    "acoustic_rating": ["Acoustic Rating"],
    "hwset": ["hwset", "HW Set", "Hardware Set", "Hardware"],
    "door_glazing": ["Door Glazing"],
    "door_glazing_thickness": ["Door Glazing (Thickness)"],
    "door_glazing_vision": ["Door Glazing and Vision Panel"],
    "comments": ["Remark", "Comments", "Description"],
    "specifier_remarks": ["specifier remarks", "Specifier Remarks"],
    "note1": ["Note 1"],
    "note2": ["Note 2"],
    "note3": ["Note 3"],
    "note4": ["Note 4"],
    "rfi": ["RFI"],
    "depth": ["Depth"],
    "arch_door_material": ["Arch Door Material", "Door Material"],
    "arch_door_type": ["Arch Door Type", "Door Type"],
    "arch_frame_material": ["Arch Frame Material", "Frame Material"],
    "arch_frame_type": ["Arch Frame Type", "Frame Type"],
    "arch_door_finish": ["Arch Door Finish"],
    "arch_frame_finish": ["Arch Frame Finish"],
    "fitout_door_model": [],
}

NUMERIC_FIELD_KEYS = {
    "width",
    "leaf1_width",
    "leaf2_width",
    "height",
    "thickness",
    "qty",
    "depth",
}


def _normalize_key(value: str) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"[\s_\-]+", " ", text)
    text = re.sub(r"[^\w\u4e00-\u9fff()]+", "", text, flags=re.UNICODE)
    return text


def _alias_lookup() -> dict[str, str]:
    """Normalized AAOS header → logical field key."""
    lookup: dict[str, str] = {}
    for field, aliases in AAOS_FIELD_ALIASES.items():
        for alias in aliases:
            lookup[_normalize_key(alias)] = field
    return lookup


_ALIAS_LOOKUP = _alias_lookup()


def country_display_name(code: str) -> str:
    return COUNTRY_NAMES.get(code.upper(), code)


def list_countries(template_root: Path | None = None) -> list[str]:
    root = template_root or TEMPLATE_ROOT
    if not root.exists():
        return []
    return sorted(
        p.name
        for p in root.iterdir()
        if p.is_dir() and any(p.glob("*.xlsx"))
    )


def list_country_templates(template_root: Path | None = None) -> list[dict[str, Any]]:
    """Return metadata for each available country template."""
    root = template_root or TEMPLATE_ROOT
    items: list[dict[str, Any]] = []
    for code in list_countries(root):
        path = find_template_path(code, root)
        try:
            columns = load_schedule_columns(path)
        except Exception:  # noqa: BLE001
            columns = []
        items.append(
            {
                "code": code,
                "name": country_display_name(code),
                "label": f"{code} - {country_display_name(code)}",
                "path": path,
                "file_name": path.name,
                "column_count": len(columns),
                "columns": columns,
            }
        )
    return items


def find_template_path(country: str, template_root: Path | None = None) -> Path:
    root = template_root or TEMPLATE_ROOT
    country_dir = root / country
    if not country_dir.exists():
        raise FileNotFoundError(f"No template folder for country '{country}'.")
    matches = sorted(country_dir.glob("*Door_Schedule*Import*Template*.xlsx"))
    if not matches:
        matches = sorted(country_dir.glob("*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"No Excel template found under {country_dir}")
    return matches[0]


def load_schedule_columns(template_path: Path) -> list[str]:
    """Read AAOS Schedule sheet header row."""
    xl = pd.ExcelFile(template_path)
    sheet = "Schedule" if "Schedule" in xl.sheet_names else xl.sheet_names[0]
    header = pd.read_excel(template_path, sheet_name=sheet, header=None, nrows=1)
    return [
        str(c).strip()
        for c in header.iloc[0].tolist()
        if c is not None and str(c).strip() not in {"", "nan", "None"}
    ]


def resolve_schedule_sheet(template_path: Path) -> str:
    xl = pd.ExcelFile(template_path)
    if "Schedule" in xl.sheet_names:
        return "Schedule"
    return xl.sheet_names[0]


def _logical_field_for_aaos_column(aaos_col: str) -> str | None:
    return _ALIAS_LOOKUP.get(_normalize_key(aaos_col))


def suggest_column_mapping(
    source_columns: list[str],
    aaos_columns: list[str],
    *,
    country: str | None = None,  # reserved for future country-specific overrides
) -> dict[str, str | None]:
    """
    Suggest mapping: AAOS column → source column name (or None).

    Uses cross-country aliases (EN/ZH), then exact/fuzzy name match.
    """
    del country  # currently unused; kept for API stability
    source_by_lower = {str(c).strip().lower(): c for c in source_columns}
    source_by_norm = {_normalize_key(c): c for c in source_columns}
    mapping: dict[str, str | None] = {}

    for aaos_col in aaos_columns:
        chosen: str | None = None
        field = _logical_field_for_aaos_column(aaos_col)
        if field:
            for candidate in SOURCE_FOR_FIELD.get(field, []):
                hit = source_by_lower.get(candidate.lower())
                if hit is not None:
                    chosen = hit
                    break
        if chosen is None:
            chosen = source_by_lower.get(aaos_col.lower())
        if chosen is None:
            chosen = source_by_norm.get(_normalize_key(aaos_col))
        mapping[aaos_col] = chosen
    return mapping


def mapping_coverage(column_mapping: dict[str, str | None]) -> dict[str, Any]:
    mapped = [k for k, v in column_mapping.items() if v]
    unmapped = [k for k, v in column_mapping.items() if not v]
    return {
        "mapped_count": len(mapped),
        "unmapped_count": len(unmapped),
        "total": len(column_mapping),
        "mapped": mapped,
        "unmapped": unmapped,
    }


def apply_aaos_mapping(
    source_df: pd.DataFrame,
    column_mapping: dict[str, str | None],
    aaos_columns: list[str],
) -> pd.DataFrame:
    """Build an AAOS Schedule dataframe from source rows using the column map."""
    if source_df is None or source_df.empty:
        raise ValueError("Complete door schedule is empty. Finish Step 2 first.")

    rows: dict[str, Any] = {}
    for aaos_col in aaos_columns:
        source_col = column_mapping.get(aaos_col)
        if source_col and source_col in source_df.columns:
            rows[aaos_col] = source_df[source_col].values
        else:
            rows[aaos_col] = [None] * len(source_df)

    result = pd.DataFrame(rows, columns=aaos_columns)

    # Coerce numeric-like AAOS columns (EN + localized headers)
    for col in result.columns:
        field = _logical_field_for_aaos_column(col)
        if field in NUMERIC_FIELD_KEYS or _normalize_key(col) in {
            "width",
            "height",
            "thickness",
            "qty",
            "quantity",
            "depth",
        }:
            result[col] = pd.to_numeric(result[col], errors="coerce")

    return result.reset_index(drop=True)


def aaos_workbook_bytes(
    aaos_df: pd.DataFrame,
    template_path: Path,
    *,
    schedule_sheet: str | None = None,
) -> bytes:
    """
    Write mapped rows into a copy of the country template workbook.

    Keeps other sheets (e.g. Data lookup lists) from the template intact.
    """
    sheet_name = schedule_sheet or resolve_schedule_sheet(template_path)
    wb = load_workbook(template_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Template is missing sheet '{sheet_name}'. Found: {wb.sheetnames}"
        )

    ws = wb[sheet_name]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    headers = [cell.value for cell in ws[1]]
    ordered_cols: list[str | None] = []
    for header in headers:
        if header is None:
            ordered_cols.append(None)
            continue
        name = str(header).strip()
        if name in aaos_df.columns:
            ordered_cols.append(name)
            continue
        match = next(
            (c for c in aaos_df.columns if _normalize_key(c) == _normalize_key(name)),
            None,
        )
        ordered_cols.append(match)

    for r_idx, (_, row) in enumerate(aaos_df.iterrows(), start=2):
        for c_idx, col_name in enumerate(ordered_cols, start=1):
            if col_name is None:
                continue
            value = row[col_name]
            if pd.isna(value):
                ws.cell(row=r_idx, column=c_idx, value=None)
            else:
                if hasattr(value, "item"):
                    try:
                        value = value.item()
                    except (ValueError, AttributeError):
                        pass
                ws.cell(row=r_idx, column=c_idx, value=value)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def mapping_table_for_editor(
    column_mapping: dict[str, str | None],
) -> pd.DataFrame:
    """Editable mapping table with logical field hints."""
    rows = []
    for aaos, source in column_mapping.items():
        field = _logical_field_for_aaos_column(aaos) or ""
        rows.append(
            {
                "AAOS Column": aaos,
                "Field": field,
                "Source Column": source or "",
            }
        )
    return pd.DataFrame(rows)


def mapping_dict_from_editor(editor_df: pd.DataFrame) -> dict[str, str | None]:
    result: dict[str, str | None] = {}
    for _, row in editor_df.iterrows():
        aaos = str(row.get("AAOS Column", "")).strip()
        source = row.get("Source Column")
        if not aaos:
            continue
        if source is None or (isinstance(source, float) and pd.isna(source)):
            result[aaos] = None
        else:
            text = str(source).strip()
            result[aaos] = text or None
    return result
